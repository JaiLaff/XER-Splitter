import os
import csv
import openpyxl as xl


class Table:
    def __init__(self, table_name:str=None, plan_name:str=None):
        self.name = table_name
        self.plan_name = plan_name
        self.rows = []
        self.writers = {
            'csv': self.WriteCSV,
            'xlsx': self.WriteXLSX
        }
        self.readers = {
            'csv': None,
            'xlsx': None
        }

    def addData(self, row:list=[]):
        if len(row) > 0:
            self.rows.append(row)

    def WriteTable(self, filetype, output_dir):
        print(f"Writing: {self.name} to {filetype} with {len(self.rows)} rows")

        write_method = self.writers.get(filetype)

        if write_method is not None:
            write_method(output_dir)
        else:
            print(f"ERROR: Internal Error, filetype {filetype} not supported")

    def WriteCSV(self,outputDir):
        try:
            with open(os.path.join(outputDir,self.name + ".csv"), "w+", newline="") as outFile:
                csv.writer(outFile, quoting=csv.QUOTE_NONNUMERIC).writerows(self.rows)

        except BaseException as e:
            print(f"Critical error during writing of table {self.name}")
            print(f"{type(e).__name__} was caught")
            print(str(e))
        else:
            print(f"INFO: {self.name} written to file successfully")


    def WriteXLSX(self,outputDir):

        write_path = os.path.join(outputDir, f'/{self.plan_name}.xlsx')
        wb = None
        sheetToWrite = None

        try:
            if not os.path.exists(write_path):
                # New excel workbook
                wb = xl.Workbook()
                sheetToWrite = wb.active
                sheetToWrite.title = self.name
            else:
                # open existing workbook
                wb = xl.load_workbook(write_path)
                sheetToWrite = wb.create_sheet(self.name)

            # The write
            for row, rowval in enumerate(self.rows, start=1):
                for col, colval in enumerate(rowval, start=1):
                    sheetToWrite.cell(row=row, column=col).value = colval
            wb.save(write_path)

        except BaseException as e:
            print(f"Critical error during writing of table {table.name}")
            print(f"{type(e).__name__} was caught")
            print(str(e))
        else:
            print(f"INFO: {table.name} written to file successfully")
    