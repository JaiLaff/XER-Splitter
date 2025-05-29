import csv
import sys
import os
import time
import argparse
import threading
import openpyxl as xl
from Plan import Plan
from Table import Table
from GUI import GUI


class Splitter:
    def __init__(self):
        self.SKIPPABLE_TABLES = ["POBS","RISKTYPE"]
        self.args = self.ParseArgs()
        self.gui = self.AttachGui()
        self.plan = None
        self.ZeroTableStats()


    def Start(self):
        if self.gui is not None: 
            self.RunGui()
        else:
            self.SplitXer()


    def ParseArgs(self):
        parser = argparse.ArgumentParser(description="A script to parse those pesky .xer files from Primavera P6", prog="xersplitter")

        fileTypeGroup = parser.add_mutually_exclusive_group()
        fileTypeGroup.add_argument("-csv", help="Comma seperated output", action="store_const", dest="type", const="csv")
        fileTypeGroup.add_argument("-xlsx", help="Excel file output", action="store_const",dest="type", const="xlsx")

        parser.add_argument("-i","--inputFile", help="The path to the input .xer file",type=str,default="", metavar="")
        parser.add_argument("-o","--outputDir", help="The directory where the output files will be placed", type=str,default="",metavar="")
        parser.add_argument("-cli", "--suppressGui", help="Hide the GUI - opens by default" ,action="store_true")
        parser.add_argument("-a", "--allTables", help="Parse all tables - Skips possibly problematic RISKTYPE & POBS tables by default", action="store_true")
        parser.add_argument("-s", "--stitch", help="Stitch all output files into a single XER file", action="store_true")

        parser.set_defaults(type="csv")
        
        args = parser.parse_args()
        return args


    def AttachGui(self):
        if not self.args.suppressGui: 
            return GUI(self.args)


    def RunGui(self):
        while True:
            event = self.gui.ProcessEvent()

            if event == 1: self.StartSplitThread()
            elif event == -1: break

        print("test")

    def ZeroTableStats(self):
        self.totalTables = 0
        self.wrTables = 0
        self.skippedTables = 0


    def UpdateStatsToGui(self):
        skippedTableText = ''
        if self.skippedTables != 0: 
            skippedTableText = f", {self.skippedTables} Tables Skipped"

        self.gui.Update(f"Written: {str(self.wrTables)}/{str(self.totalTables)} Tables{skippedTableText}")


    def CheckDirectories(self):
        if not os.path.exists(self.args.inputFile):
            print(f"ERROR: Could not find target xer file \"{self.args.inputFile}\"")
            return False

        if not os.path.exists(self.args.outputDir):
            print(f"INFO: Could not find target output directory \"{self.args.outputDir}\"")
            print(f"Attempting to create target output directory...")
            
            try:
                os.makedirs(self.args.outputDir)
            except BaseException as e:
                print(f"ERROR: Failed to create target output directory \"{self.args.outputDir}\"")
                print(f"{type(e).__name__} was caught")
                print(str(e))
                return False
            else: print(f"Target output directory \"{self.args.outputDir}\" created successfully.")
        
        return True


    def PreCheck(self):
        print("INFO: Beginning PreCheck of the XER File")
        
        eof = False

        try:
            with open(self.args.inputFile, "r", encoding="cp1252", errors="ignore") as xer:
                

                while not eof:
                    rowType = xer.readline(2)
                    if rowType == "%T": self.totalTables += 1
                    elif rowType == "%E": eof = True
                    
                self.UpdateStatsToGui()

        except BaseException as e:
            print("Critical error during Pre Check of XER File")
            print(f"{type(e).__name__} was caught")
            print(str(e))
            return False

        print("INFO: Pre Check completed successfully")
        print(f"{self.totalTables} tables found")
        return True


    def WriteCSV(self,outputDir, table:Table):
        try:
            with open(os.path.join(outputDir,table.name + ".csv"), "w+", newline="") as outFile:

                csv.writer(outFile, quoting=csv.QUOTE_NONNUMERIC).writerows(table.rows)
        except BaseException as e:
            print(f"Critical error during writing of table {table.name}")
            print(f"{type(e).__name__} was caught")
            print(str(e))
        else:
            print(f"INFO: {table.name} written to file successfully")


    def WriteXLSX(self,outputDir, table:Table):
        
    
        try:
            wb = None
            sheetToWrite = None
            if not os.path.exists(outputDir):
                # New excel workbook
                wb = xl.Workbook()
                sheetToWrite = wb.active
                sheetToWrite.title = table.name
            else:
                # open existing workbook
                wb = xl.load_workbook(outputDir)
                sheetToWrite = wb.create_sheet(table.name)

            # The write
            for row, rowval in enumerate(table.rows, start=1):
                for col, colval in enumerate(rowval, start=1):
                    sheetToWrite.cell(row=row, column=col).value = colval
            wb.save(outputDir)

        except BaseException as e:
            print(f"Critical error during writing of table {table.name}")
            print(f"{type(e).__name__} was caught")
            print(str(e))
        else:
            print(f"INFO: {table.name} written to file successfully")


    def WriteTable(self, table:Table):
        print(f"Writing: {table.name} with {len(table.rows)} rows")

        if self.args.type == "csv":
            self.WriteCSV(self.args.outputDir,table)
        elif self.args.type == "xlsx":
            # constructing output file dir
            tail = os.path.split(self.args.inputFile)[1]
            filename = tail[:-3]
            outputDir = os.path.join(self.args.outputDir, filename)
            self.WriteXLSX(outputDir + "xlsx",table)
        
        self.wrTables += 1
        self.UpdateStatsToGui()


    def Split(self):
        eof = False
        plan = Plan(plan_name=self.args.inputFile)
        tableToWrite = None

        try:
            with open(self.args.inputFile, 'r', encoding="cp1252", errors="ignore") as xer:
                while not eof:
                    currentLine = xer.readline().split('\t')
                    currentLine[len(currentLine)-1] = currentLine[len(currentLine)-1][:-1]

                    rowType = currentLine[0]

                    if rowType == "%E":
                        eof = True

                    if rowType == "%T":

                        # Write previous table if available and flush the rows
                        if tableToWrite:
                            plan.addTable(tableToWrite)
                            self.WriteTable(tableToWrite)
                            tableToWrite = None

                        tableToWrite = Table(table_name=currentLine[1])

                        if not self.args.allTables:
                            if tableToWrite.name in self.SKIPPABLE_TABLES:
                                print(f"Found table \"{tableToWrite.name}\" - Skipping")
                                self.skippedTables += 1

                                while xer.readline(2) != "%T":
                                    # save line position for when a table is found
                                    savedLine = xer.tell()

                                # new table found -> go back to the line and read the whole line
                                xer.seek(savedLine)
                                continue

                        print(f"Found table \"{tableToWrite.name}\"")

                    if rowType in ["%R", "%F"]:
                        currentLine.pop(0)
                        tableToWrite.addData(currentLine)

                # Write remaining tables
                if tableToWrite:
                    plan.addTable(tableToWrite)
                    self.WriteTable(tableToWrite)
                    tableToWrite=None

        except BaseException as e:
            print(f"Critical error splitting XER")
            print(f"{type(e).__name__} was caught")
            print(str(e))
            raise
        else:
            print(f"INFO: Split complete successfully")
            return True
        

    def SplitXer(self):
        print("INFO: Settings Confirmed")
        print(self.args)
        self.ZeroTableStats()

        if not (self.CheckDirectories() and self.PreCheck() and self.Split()): 
            self.gui.Update(success=False)
            return

        self.gui.Update(success=True)
        

    def StartSplitThread(self):
        print("INFO: Starting Compute Thread...")
        threading.Thread(target=self.SplitXer, daemon=True).start()


if __name__ == "__main__":
    splitter = Splitter()
    splitter.Start()